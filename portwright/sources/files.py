"""File adapter: CSV, JSON (array of objects / JSON-lines) and Excel (.xlsx).

Security properties:
* Files must live inside ``security.data_root`` – path traversal and symlink
  escapes are rejected at build time AND at read time.
* Read-only. Only declared ``filters`` can be filtered on; ``columns`` acts as
  a field allow-list so sensitive columns are never exposed.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import anyio

from ..config import EndpointSpec
from .base import DataSource, ExecutionContext, NotFound, SourceError, register

_SUPPORTED = {".csv", ".json", ".jsonl", ".xlsx", ".xlsm"}
_MAX_FILE_BYTES = 200 * 1024 * 1024


def decode_text(data: bytes, encoding: str = "auto") -> str:
    """Bytes -> text: a BOM wins, then UTF-8, then Windows-1252 (encoding='auto'), or the given codec."""
    if encoding and encoding != "auto":
        return data.decode(encoding)
    for bom, codec in ((b"\xef\xbb\xbf", "utf-8-sig"), (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16")):
        if data.startswith(bom):
            return data.decode(codec)
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def parse_file(path: Path, sheet: str | None = None, encoding: str = "auto") -> list[dict[str, Any]]:
    """Rows of a CSV, JSON (array or {items: [...]}), JSON Lines or Excel file."""
    suffix = path.suffix.lower()
    if suffix not in _SUPPORTED:
        raise SourceError(f"cannot read rows from a '{suffix}' file (CSV, JSON, JSONL or Excel only)", 500)
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.active
            it = ws.iter_rows(values_only=True)
            header = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(next(it, []))]
            return [dict(zip(header, r)) for r in it if any(v is not None for v in r)]
        finally:
            wb.close()
    text = decode_text(path.read_bytes(), encoding)
    if suffix == ".csv":
        import io
        return list(csv.DictReader(io.StringIO(text, newline="")))
    if suffix == ".jsonl":
        return [json.loads(line) for line in text.splitlines() if line.strip()]
    data = json.loads(text)
    return data if isinstance(data, list) else data.get("items", [])


@register("file")
class FileSource(DataSource):
    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self._cache: tuple[float, list[dict[str, Any]]] | None = None

    # ------------------------------------------------------------- sandboxing
    def _resolved_path(self) -> Path:
        try:
            root = (Path(self.project_root) / self.security.data_root).resolve()
            target = (root / self.spec.path).resolve()          # resolves symlinks too
        except OSError as exc:                                  # e.g. WinError 123 (bad characters)
            raise SourceError(f"Invalid path for source '{self.name}': path={self.spec.path!r}, "
                              f"data_root={self.security.data_root!r} ({exc.strerror or exc}). "
                              f"Use just the file name inside the data folder, e.g. 'transfer.csv'",
                              500) from None
        if root != target and root not in target.parents:
            raise SourceError(f"File for source '{self.name}' ({self.spec.path!r}) is outside the data "
                              f"folder; copy it into '{self.security.data_root}' and use just its name", 500)
        if target.suffix.lower() not in _SUPPORTED:
            raise SourceError(f"Unsupported file type '{target.suffix}'", 500)
        return target

    def validate_endpoint(self, ep: EndpointSpec) -> None:
        if ep.method != "GET":
            raise ValueError(f"[{ep.name}] file sources are read-only (GET only)")
        try:
            self._resolved_path()
        except SourceError as exc:
            raise ValueError(f"[{ep.name}] {exc}") from None
        if ep.key and not any(p.name == ep.key and p.location == "path" for p in ep.params):
            raise ValueError(f"[{ep.name}] key '{ep.key}' must be a path param")

    async def health(self) -> bool:
        try:
            return self._resolved_path().exists()
        except SourceError:
            return False

    # ---------------------------------------------------------------- loading
    def _load(self) -> list[dict[str, Any]]:
        path = self._resolved_path()
        if not path.exists():
            raise SourceError("Data file not found", 500)
        stat = path.stat()
        if stat.st_size > _MAX_FILE_BYTES:
            raise SourceError("Data file too large", 500)
        if self._cache and self._cache[0] == stat.st_mtime:
            return self._cache[1]

        rows = parse_file(path, self.spec.sheet)
        self._cache = (stat.st_mtime, rows)
        return rows

    # ---------------------------------------------------------------- execute
    async def execute(self, ctx: ExecutionContext) -> Any:
        rows = await anyio.to_thread.run_sync(self._load)
        ep = ctx.endpoint

        def project(r: dict[str, Any]) -> dict[str, Any]:
            return {c: r.get(c) for c in ep.columns} if ep.columns else dict(r)

        if ep.key:  # single-record lookup
            wanted = str(ctx.params.get(ep.key))
            for r in rows:
                if str(r.get(ep.key)) == wanted:
                    return project(r)
            raise NotFound()

        active = {f: str(ctx.params[f]) for f in ep.filters
                  if ctx.params.get(f) is not None}
        matched = [r for r in rows if all(str(r.get(k)) == v for k, v in active.items())]
        page = matched[ctx.offset: ctx.offset + ctx.limit]
        return {"items": [project(r) for r in page], "limit": ctx.limit,
                "offset": ctx.offset, "count": len(page), "total": len(matched)}
